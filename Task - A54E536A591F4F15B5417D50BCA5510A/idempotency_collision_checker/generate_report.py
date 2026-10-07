import os
import pandas as pd

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUTPUT_FOLDER = "output"

CSV_FILE = os.path.join(
    OUTPUT_FOLDER,
    "retry_analysis.csv"
)

COLLISION_FILE = os.path.join(
    OUTPUT_FOLDER,
    "collision_details.csv"
)

GRAPH_FILE = os.path.join(
    OUTPUT_FOLDER,
    "retry_summary.png"
)

WORD_FILE = os.path.join(
    OUTPUT_FOLDER,
    "reliability_note.docx"
)


# ==========================================================
# DOCUMENT FORMATTING
# ==========================================================

def set_cell_shading(cell, fill):

    tc_pr = cell._tc.get_or_add_tcPr()

    shd = OxmlElement("w:shd")

    shd.set(
        qn("w:fill"),
        fill
    )

    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):

    cell.text = ""

    paragraph = cell.paragraphs[0]

    run = paragraph.add_run(
        str(text)
    )

    run.bold = bold
    run.font.size = Pt(8.5)

    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table):

    tbl = table._tbl

    tbl_pr = tbl.tblPr

    borders = OxmlElement("w:tblBorders")

    for edge in (
        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV"
    ):

        tag = "w:" + edge

        element = OxmlElement(tag)

        element.set(
            qn("w:val"),
            "single"
        )

        element.set(
            qn("w:sz"),
            "4"
        )

        element.set(
            qn("w:space"),
            "0"
        )

        element.set(
            qn("w:color"),
            "B7B7B7"
        )

        borders.append(element)

    tbl_pr.append(borders)


def add_heading(document, text, level=1):

    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(3)

    run = paragraph.add_run(text)

    run.bold = True

    if level == 1:
        run.font.size = Pt(11)

    else:
        run.font.size = Pt(10)

    return paragraph


def add_body(document, text):

    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.05

    run = paragraph.add_run(text)

    run.font.size = Pt(9)

    return paragraph


# ==========================================================
# LOAD RESULTS
# ==========================================================

results = pd.read_csv(
    CSV_FILE
)

collisions = pd.read_csv(
    COLLISION_FILE
)


# ==========================================================
# CALCULATE SUMMARY
# ==========================================================

first_use_count = (
    results["status"] == "FIRST_USE"
).sum()

safe_retry_count = (
    results["status"] == "SAFE_RETRY"
).sum()

collision_count = (
    results["status"] == "COLLISION"
).sum()

trace_event_count = (
    results["status"] == "TRACE_EVENT"
).sum()


# ==========================================================
# CREATE DOCUMENT
# ==========================================================

document = Document()

section = document.sections[0]

section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.6)
section.right_margin = Inches(0.6)


# Default font

styles = document.styles

styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(9)


# ==========================================================
# TITLE
# ==========================================================

title = document.add_paragraph()

title.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = title.add_run(
    "IDEMPOTENCY-KEY COLLISION CHECK"
)

run.bold = True
run.font.size = Pt(16)

subtitle = document.add_paragraph()

subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = subtitle.add_run(
    "Reliability Test Note — Fictional API Retry Trace"
)

run.font.size = Pt(9)


# ==========================================================
# EXECUTIVE SUMMARY
# ==========================================================

add_heading(
    document,
    "1. Executive Summary"
)

add_body(
    document,
    "A small Python checker was used to review a fictional API "
    "retry trace for unsafe reuse of idempotency keys. The test "
    "successfully identified one legitimate retry where the same "
    "key was reused with the same request payload, and one deliberate "
    "collision where the same key was reused with different request data."
)

add_body(
    document,
    "Overall result: REVIEW REQUIRED. The collision condition was "
    "detected as intended and should be investigated before treating "
    "the trace as reliability-safe."
)


# ==========================================================
# WHAT WAS TESTED
# ==========================================================

add_heading(
    document,
    "2. What Was Tested"
)

add_body(
    document,
    "The test contained eight fictional API events covering an "
    "initial request, a timeout, a valid retry, normal request "
    "activity and a deliberate same-key/different-payload collision."
)

add_body(
    document,
    "No live API calls, production logs or production data were used. "
    "The purpose was to verify that the checker can distinguish a "
    "safe retry from unsafe reuse of an idempotency key."
)


# ==========================================================
# HOW THE CHECKER WORKS
# ==========================================================

add_heading(
    document,
    "3. How the Checker Works"
)

add_body(
    document,
    "The checker records the request payload associated with each "
    "idempotency key. When the same key appears again as a request "
    "or retry, the new payload is compared with the original."
)

add_body(
    document,
    "Same key + same payload = SAFE RETRY. "
    "Same key + different payload = COLLISION. "
    "Timeout and response events are retained for trace context "
    "but are not treated as new requests."
)


# ==========================================================
# RESULTS TABLE
# ==========================================================

add_heading(
    document,
    "4. Event-by-Event Results"
)

table = document.add_table(
    rows=1,
    cols=5
)

table.alignment = WD_TABLE_ALIGNMENT.CENTER

table.style = "Table Grid"

headers = [
    "Event",
    "Type",
    "Key",
    "Finding",
    "Interpretation"
]

for index, header in enumerate(headers):

    cell = table.rows[0].cells[index]

    set_cell_text(
        cell,
        header,
        bold=True
    )

    set_cell_shading(
        cell,
        "D9EAF7"
    )


for _, row in results.iterrows():

    cells = table.add_row().cells

    status = row["status"]

    if status == "FIRST_USE":

        interpretation = "New key"

    elif status == "SAFE_RETRY":

        interpretation = "Same key + same payload"

    elif status == "COLLISION":

        interpretation = "Same key + different payload"

    else:

        interpretation = "Trace context only"

    values = [
        row["event_id"],
        row["event"],
        row["key"],
        status,
        interpretation
    ]

    for index, value in enumerate(values):

        set_cell_text(
            cells[index],
            value
        )

    if status == "COLLISION":

        for cell in cells:

            set_cell_shading(
                cell,
                "FCE4D6"
            )


set_table_borders(table)


# ==========================================================
# SUMMARY VISUALISATION
# ==========================================================

add_heading(
    document,
    "5. Test Findings at a Glance"
)

if os.path.exists(GRAPH_FILE):

    paragraph = document.add_paragraph()

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run()

    run.add_picture(
        GRAPH_FILE,
        width=Inches(5.6)
    )


add_body(
    document,
    f"The checker classified {first_use_count} event(s) as first use, "
    f"{safe_retry_count} as safe retries, and {collision_count} as "
    f"collision(s). {trace_event_count} additional events were retained "
    f"as timeout/response trace context."
)


# ==========================================================
# COLLISION DETAIL
# ==========================================================

add_heading(
    document,
    "6. Collision Detail"
)

if not collisions.empty:

    for _, row in collisions.iterrows():

        add_body(
            document,
            f"Key '{row['key']}' was originally associated with event "
            f"{row['original_event_id']}. At event {row['event_id']}, "
            f"the same key was reused with different request data."
        )

        collision_table = document.add_table(
            rows=1,
            cols=4
        )

        collision_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        collision_table.style = "Table Grid"

        collision_headers = [
            "Idempotency key",
            "Original value",
            "New value",
            "Changed field"
        ]

        for index, header in enumerate(collision_headers):

            cell = collision_table.rows[0].cells[index]

            set_cell_text(
                cell,
                header,
                bold=True
            )

            set_cell_shading(
                cell,
                "F4CCCC"
            )

        cells = collision_table.add_row().cells

        values = [
            row["key"],
            row["original_value"],
            row["new_value"],
            row["changed_field"]
        ]

        for index, value in enumerate(values):

            set_cell_text(
                cells[index],
                value
            )

        set_table_borders(
            collision_table
        )


# ==========================================================
# WHY IT MATTERS
# ==========================================================

add_heading(
    document,
    "7. Why This Matters"
)

add_body(
    document,
    "An idempotency key is intended to identify one logical operation "
    "across retries. Reusing that key with different request data "
    "creates ambiguity about which operation the key represents."
)

add_body(
    document,
    "In a real API environment, this could result in rejected "
    "requests, incorrect deduplication or inconsistent application "
    "behaviour, depending on how the API implements idempotency. "
    "The checker therefore treats same-key/different-payload reuse "
    "as a reliability risk rather than assuming that it is a normal retry."
)


# ==========================================================
# RECOMMENDATION
# ==========================================================

add_heading(
    document,
    "8. Recommendation"
)

add_body(
    document,
    "REVIEW REQUIRED. The prototype correctly detected the deliberate "
    "collision. The affected key and request sequence should be "
    "investigated before considering the behaviour safe."
)

add_body(
    document,
    "The checker is suitable as a local trace-analysis prototype. "
    "Before production use, it should be tested against additional "
    "cases such as multiple retries, missing keys, malformed payloads, "
    "different field ordering and repeated collisions."
)


# ==========================================================
# SCOPE
# ==========================================================

add_heading(
    document,
    "9. Scope and Limitations"
)

add_body(
    document,
    "This exercise analyses a small fictional trace locally. It does "
    "not call an API, modify application data or validate the behaviour "
    "of a real production idempotency implementation. The findings "
    "demonstrate the checker's detection logic rather than a production incident."
)


# ==========================================================
# SAVE DOCUMENT
# ==========================================================

document.save(
    WORD_FILE
)

print()
print("=" * 70)
print("WORD REPORT GENERATED")
print("=" * 70)
print(
    f"File: {WORD_FILE}"
)
print(
    f"First uses: {first_use_count}"
)
print(
    f"Safe retries: {safe_retry_count}"
)
print(
    f"Collisions: {collision_count}"
)
print("=" * 70)